#!/usr/bin/env python
"""
eval/run_eval.py
────────────────
Generic evaluation runner for the Veltrix PeopleOps Assistant.

It reads a question set (questions + expected answers + grading rules) from an
EXTERNAL data file, calls the running API once per question, records the actual
answer, auto-grades it, and writes everything to an Excel workbook with a summary
sheet (success percentage).

IMPORTANT — no document content is hardcoded in this script. Every fact, expected
answer and grading rule lives in the dataset file (default:
eval/dataset/qa_dataset.json). To evaluate a modified corpus, point --dataset at a
new file; the code does not change.

Usage
-----
    # 1. start the API (with an LLM key in .env)
    uvicorn src.main:app --port 8000

    # 2. run the evaluation
    python eval/run_eval.py

    # build the Excel skeleton without calling the API
    python eval/run_eval.py --dry-run

    # evaluate a subset
    python eval/run_eval.py --only Q05,Q06,Q14
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import urllib.error
import urllib.request
import uuid
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))

DEFAULT_DATASET = os.path.join(HERE, "dataset", "qa_dataset.json")
DEFAULT_OUT = os.path.join(HERE, "results", "eval_results.xlsx")
DEFAULT_BASE_URL = os.environ.get("EVAL_BASE_URL", "http://localhost:8000")


# ─── API client ───────────────────────────────────────────────────────────────

def call_chat(base_url: str, employee_id: str, session_id: str, message: str,
              timeout: float = 90.0) -> dict:
    """POST /chat. Returns the parsed JSON response, or {error: ...} on failure."""
    payload = json.dumps({"session_id": session_id, "message": message}).encode("utf-8")
    req = urllib.request.Request(
        f"{base_url.rstrip('/')}/chat",
        data=payload,
        headers={"Content-Type": "application/json", "X-Employee-Id": employee_id},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", "replace")
        return {"error": f"HTTP {exc.code}: {body}"}
    except Exception as exc:  # noqa: BLE001
        return {"error": f"{type(exc).__name__}: {exc}"}


def check_health(base_url: str, timeout: float = 10.0) -> tuple[bool, str]:
    try:
        with urllib.request.urlopen(f"{base_url.rstrip('/')}/health", timeout=timeout) as resp:
            return resp.status == 200, resp.read().decode("utf-8", "replace")
    except urllib.error.HTTPError as exc:
        return False, f"HTTP {exc.code}"
    except Exception as exc:  # noqa: BLE001
        return False, f"{type(exc).__name__}: {exc}"


def call_with_retries(base_url: str, employee_id: str, session_id: str, message: str,
                      timeout: float, retries: int, retry_wait: float,
                      label: str = "") -> tuple[dict, str, int, int]:
    """
    POST /chat, retrying transient failures.

    A response is treated as transient when the HTTP call failed or when the API
    reports ``degraded: true`` (the assistant could not reach the LLM provider,
    e.g. the provider rate-limited us). Each retry uses a FRESH session id so the
    failed turn does not pollute the conversation history.

    Returns (response, session_id, retries_used, latency_ms_of_last_attempt).
    """
    attempt = 0
    while True:
        t0 = time.time()
        resp = call_chat(base_url, employee_id, session_id, message, timeout)
        latency_ms = int((time.time() - t0) * 1000)
        healthy = "error" not in resp and not resp.get("degraded")
        if healthy or attempt >= retries:
            return resp, session_id, attempt, latency_ms

        attempt += 1
        wait = retry_wait * attempt
        reason = "http-error" if "error" in resp else "degraded (provider unavailable)"
        print(f"      retry {attempt}/{retries} for {label} in {wait:.0f}s — {reason}")
        time.sleep(wait)
        session_id = f"{session_id.rsplit('-', 1)[0]}-{uuid.uuid4().hex[:8]}"


# ─── Grading (rules come from the dataset, never hardcoded here) ──────────────

# LLM output often uses typographic characters (narrow no-break space U+202F,
# non-breaking space U+00A0, non-breaking hyphen U+2011, en/em dashes) that a
# human reads as a plain space/hyphen. Normalising before matching means a
# correct answer is never marked wrong just for its punctuation style.
_TYPOGRAPHIC = {
    "\u202f": " ", "\u00a0": " ", "\u2009": " ", "\u2007": " ", "\u2060": "",
    "\u2011": "-", "\u2010": "-", "\u2012": "-", "\u2013": "-", "\u2014": "-",
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
}


def normalise_for_match(text: str) -> str:
    """Fold typographic whitespace/dashes and collapse runs of whitespace."""
    out = (text or "")
    for src, dst in _TYPOGRAPHIC.items():
        out = out.replace(src, dst)
    return " ".join(out.split())


def grade(answer: str, g: dict) -> dict:
    low = normalise_for_match(answer).lower()

    missing = [s for s in g.get("must_contain", []) if normalise_for_match(s).lower() not in low]

    any_terms = g.get("must_contain_any", [])
    any_hits = [s for s in any_terms if normalise_for_match(s).lower() in low]
    any_ok = (not any_terms) or bool(any_hits)

    forbidden = [s for s in g.get("must_not_contain", []) if normalise_for_match(s).lower() in low]

    rx = g.get("expect_regex")
    regex_ok = True if not rx else (re.search(rx, normalise_for_match(answer), re.IGNORECASE) is not None)

    passed = (not missing) and any_ok and (not forbidden) and regex_ok
    return {
        "passed": passed,
        "missing": missing,
        "any_hits": any_hits,
        "any_ok": any_ok,
        "forbidden_hits": forbidden,
        "regex_ok": regex_ok,
    }


# ─── Excel output ─────────────────────────────────────────────────────────────

COLUMNS = [
    ("id", 8), ("category", 22), ("employee_id", 12), ("role_context", 26),
    ("question", 52), ("expected_answer", 60), ("actual_answer", 70),
    ("route", 14), ("pending_confirmation", 18), ("citations", 30), ("tool_calls", 30),
    ("auto_result", 12), ("missing_keywords", 22), ("forbidden_hits", 22),
    ("regex_ok", 9), ("review_rubric", 46),
    ("reviewer_verdict", 16), ("reviewer_notes", 30),
    ("latency_ms", 11), ("error", 30),
    # Appended last so the Summary sheet's column references stay stable.
    ("degraded", 11), ("retries", 9),
]


def write_excel(rows: list[dict], out_path: str, meta: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # ── Results sheet ──
    ws = wb.active
    ws.title = "Results"
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(color="FFFFFF", bold=True)
    for c, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A2"

    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    auto_col = next(i for i, (name, _w) in enumerate(COLUMNS, start=1) if name == "auto_result")
    for r, row in enumerate(rows, start=2):
        for c, (name, _w) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=row.get(name, ""))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        verdict = ws.cell(row=r, column=auto_col)
        verdict.fill = pass_fill if row.get("auto_result") == "PASS" else fail_fill
        ws.row_dimensions[r].height = 60

    # ── Summary sheet ──
    s = wb.create_sheet("Summary")
    total = len(rows)
    last = total + 1

    def put(r, label, value):
        s.cell(row=r, column=1, value=label).font = Font(bold=True)
        s.cell(row=r, column=2, value=value)

    put(1, "Generated", datetime.now().isoformat(timespec="seconds"))
    put(2, "Reference date", meta.get("reference_date", ""))
    put(3, "Dataset", meta.get("name", ""))
    put(4, "Base URL", meta.get("base_url", ""))
    put(5, "Total questions", total)
    put(6, "Auto-graded PASS", f'=COUNTIF(Results!L2:L{last},"PASS")')
    put(7, "Auto-graded FAIL", f'=COUNTIF(Results!L2:L{last},"FAIL")')
    put(8, "Auto-graded ERROR", f'=COUNTIF(Results!L2:L{last},"ERROR")')
    s.cell(row=9, column=1, value="Auto success rate (of graded)").font = Font(bold=True)
    s.cell(row=9, column=2, value='=IF(B6+B7=0,"",B6/(B6+B7))').number_format = "0.0%"
    # reviewer columns (filled by hand after review)
    s.cell(row=11, column=1, value="Manual review (fill 'PASS'/'FAIL' in Results!Q)").font = Font(bold=True, italic=True)
    put(12, "Manual PASS count", f'=COUNTIF(Results!Q2:Q{last},"PASS")')
    put(13, "Manual FAIL count", f'=COUNTIF(Results!Q2:Q{last},"FAIL")')
    put(14, "Manual reviewed", f'=COUNTIF(Results!Q2:Q{last},"PASS")+COUNTIF(Results!Q2:Q{last},"FAIL")')
    s.cell(row=15, column=1, value="Manual success rate (of reviewed)").font = Font(bold=True)
    s.cell(row=15, column=2, value='=IF(B14=0,"",B12/B14)').number_format = "0.0%"
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 34

    # ── By-category breakdown ──
    cats: dict[str, list[dict]] = {}
    for r in rows:
        cats.setdefault(r.get("category", "?"), []).append(r)
    s.cell(row=17, column=1, value="Category").font = Font(bold=True)
    s.cell(row=17, column=2, value="Auto pass / total").font = Font(bold=True)
    for i, (cat, items) in enumerate(sorted(cats.items()), start=18):
        p = sum(1 for it in items if it.get("auto_result") == "PASS")
        s.cell(row=i, column=1, value=cat)
        s.cell(row=i, column=2, value=f"{p} / {len(items)}")

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)


# ─── Main ─────────────────────────────────────────────────────────────────────

def main() -> int:
    ap = argparse.ArgumentParser(description="Run the Veltrix PeopleOps evaluation set.")
    ap.add_argument("--dataset", default=DEFAULT_DATASET)
    ap.add_argument("--out", default=DEFAULT_OUT)
    ap.add_argument("--base-url", default=DEFAULT_BASE_URL)
    ap.add_argument("--only", default="", help="Comma-separated question IDs to run")
    ap.add_argument("--timeout", type=float, default=90.0)
    ap.add_argument("--delay", type=float, default=5.0,
                    help="Seconds to wait between questions. Increase this if your "
                         "LLM provider has a low tokens-per-minute limit.")
    ap.add_argument("--retries", type=int, default=3,
                    help="Retries per question when the API is degraded or errors.")
    ap.add_argument("--retry-wait", type=float, default=30.0,
                    help="Base backoff for retries (multiplied by the attempt number).")
    ap.add_argument("--dry-run", action="store_true",
                    help="Write the Excel skeleton without calling the API.")
    args = ap.parse_args()

    with open(args.dataset, encoding="utf-8") as f:
        dataset = json.load(f)
    meta = dict(dataset.get("meta", {}))
    meta["base_url"] = args.base_url
    questions = dataset["questions"]

    if args.only:
        want = {x.strip() for x in args.only.split(",") if x.strip()}
        questions = [q for q in questions if q["id"] in want]

    print(f"Dataset : {args.dataset}  ({len(questions)} questions)")
    print(f"Base URL: {args.base_url}")

    if not args.dry_run:
        ok, body = check_health(args.base_url)
        print(f"Health  : {'OK' if ok else 'NOT READY'} — {body[:120]}")
        if not ok:
            print("\nThe API is not healthy. Start it (uvicorn src.main:app --port 8000) "
                  "and make sure an LLM key is set in .env.", file=sys.stderr)
            return 1

    rows: list[dict] = []
    for i, q in enumerate(questions, start=1):
        row = {
            "id": q["id"], "category": q.get("category", ""),
            "employee_id": q["employee_id"], "role_context": q.get("role_context", ""),
            "question": q["question"], "expected_answer": q.get("expected_answer", ""),
            "review_rubric": q.get("review_rubric", ""),
            "actual_answer": "", "route": "", "pending_confirmation": "",
            "citations": "", "tool_calls": "", "auto_result": "", "missing_keywords": "",
            "forbidden_hits": "", "regex_ok": "", "latency_ms": "", "error": "",
            "degraded": "", "retries": 0,
            "reviewer_verdict": "", "reviewer_notes": "",
        }

        if args.dry_run:
            row["auto_result"] = "(dry-run)"
            rows.append(row)
            continue

        session_id = f"eval-{q['id']}-{uuid.uuid4().hex[:8]}"
        resp, session_id, retries_used, latency_ms = call_with_retries(
            args.base_url, q["employee_id"], session_id, q["question"],
            args.timeout, args.retries, args.retry_wait, label=q["id"],
        )
        row["latency_ms"] = latency_ms
        row["retries"] = retries_used

        if "error" in resp:
            row["error"] = resp["error"]
            row["auto_result"] = "ERROR"
        elif resp.get("degraded"):
            # The assistant never reached the model, so there is nothing to
            # grade. Reporting this as ERROR (not FAIL) keeps the accuracy
            # figure honest instead of counting an outage as a wrong answer.
            row["error"] = "degraded: LLM provider unavailable after retries"
            row["actual_answer"] = resp.get("answer", "") or ""
            row["route"] = resp.get("route", "")
            row["degraded"] = True
            row["auto_result"] = "ERROR"
        else:
            answer = resp.get("answer", "") or ""
            # include any follow-up turns in the graded text
            for fu in q.get("followups", []):
                sid2 = f"{session_id}-fu-{uuid.uuid4().hex[:6]}"
                r2, _sid2, _rt, _lat = call_with_retries(
                    args.base_url, q["employee_id"], sid2, fu,
                    args.timeout, args.retries, args.retry_wait, label=q["id"],
                )
                if "error" not in r2 and not r2.get("degraded"):
                    answer += "\n" + (r2.get("answer", "") or "")
            row["actual_answer"] = answer
            row["route"] = resp.get("route", "")
            row["pending_confirmation"] = resp.get("pending_confirmation", "")
            row["degraded"] = resp.get("degraded", False)
            row["citations"] = "; ".join(
                f"{c.get('doc_id')}#{c.get('locator')}" for c in resp.get("citations", []))
            row["tool_calls"] = "; ".join(
                f"{t.get('name')}({t.get('status')})" for t in resp.get("tool_calls", []))
            res = grade(answer, q.get("grading", {}))
            row["auto_result"] = "PASS" if res["passed"] else "FAIL"
            row["missing_keywords"] = ", ".join(res["missing"])
            row["forbidden_hits"] = ", ".join(res["forbidden_hits"])
            row["regex_ok"] = res["regex_ok"]

        rows.append(row)
        print(f"[{i:2}/{len(questions)}] {q['id']:>4} {row['auto_result']:<6} "
              f"{row['latency_ms']}ms  {q['question'][:58]}")

        if args.delay > 0 and i < len(questions):
            time.sleep(args.delay)

    write_excel(rows, args.out, meta)

    if args.dry_run:
        print(f"\nDry-run: wrote skeleton to {args.out}")
        return 0

    total = len(rows)
    passed = sum(1 for r in rows if r["auto_result"] == "PASS")
    errored = sum(1 for r in rows if r["auto_result"] == "ERROR")
    graded = total - errored
    rate = (passed / graded * 100) if graded else 0.0
    print("\n" + "=" * 60)
    print(f"Auto-graded: {passed}/{graded} PASS  ({rate:.1f}% of graded)   "
          f"errors: {errored}/{total}")
    print(f"Excel written: {args.out}")
    print("Review each row, then set 'reviewer_verdict' (PASS/FAIL) in the Results sheet;")
    print("the Summary sheet computes the reviewed success rate automatically.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
