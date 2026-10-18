"""
src/ingestion/holiday_extractor.py
────────────────────────────────────
Extract holiday dates from calendar documents and persist to the holidays SQLite table.

Uses generic column-header detection — never hardcoded column indices.
Handles: HTML tables, DOCX tables, PDF tables.
"""
from __future__ import annotations
import os
import re
from datetime import datetime
from typing import Optional

from src.data_access.db import get_db
from src.ingestion.text_cleaner import clean_table_cell
from src.utils.logging import get_logger

logger = get_logger(__name__)

# ─── Holiday table creation ───────────────────────────────────────────────────

HOLIDAYS_TABLE_DDL = """
CREATE TABLE IF NOT EXISTS holidays (
    region       TEXT NOT NULL,
    date         TEXT NOT NULL,
    name         TEXT NOT NULL,
    year         INTEGER NOT NULL,
    doc_id       TEXT NOT NULL,
    PRIMARY KEY (region, date)
);
"""


def ensure_holidays_table() -> None:
    conn = get_db()
    conn.execute(HOLIDAYS_TABLE_DDL)
    conn.commit()


# ─── Date normalization ───────────────────────────────────────────────────────

_DATE_FORMATS = [
    "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y",
    "%d %b %Y", "%d %B %Y", "%B %d, %Y", "%b %d, %Y",
    "%d %b", "%d %B",  # No year — will use default
]


def normalize_date(raw: str, default_year: int = 2026) -> Optional[str]:
    """Normalize a date string to YYYY-MM-DD. Returns None if parsing fails."""
    raw = re.sub(r'\s+', ' ', raw.strip())
    raw = raw.rstrip('.')

    for fmt in _DATE_FORMATS:
        try:
            d = datetime.strptime(raw, fmt)
            if d.year == 1900:  # strptime default when no year in format
                d = d.replace(year=default_year)
            return d.strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


# ─── Generic table row parser ─────────────────────────────────────────────────

# Header keywords used for generic column detection (no hardcoded indices).
_DATE_STRONG = ("date", "when", "observance", "observed")
_DATE_WEAK = ("day",)                       # "Day" is often the weekday, not the date
_NAME_KEYS = ("holiday", "occasion", "event", "name", "description", "festival")
# A table whose name column is "restricted/optional holiday" lists days the office
# stays OPEN — those must NOT be subtracted from working days.
_NON_WORKING_HEADER = ("restricted", "optional", "floating")


def parse_holiday_table(rows: list[list[str]], doc_id: str, region: str) -> list[dict]:
    """
    Parse ONE holiday table using column-header detection (no hardcoded indices).

    Returns list of {region, date, name, year, doc_id} dicts. Returns [] for any
    table that is not a holiday table (e.g. a document-metadata table).
    """
    if not rows:
        return []

    headers = [clean_table_cell(h).lower() for h in rows[0]]

    date_col = next((i for i, h in enumerate(headers)
                     if any(k in h for k in _DATE_STRONG)), None)
    if date_col is None:
        date_col = next((i for i, h in enumerate(headers)
                         if any(k in h for k in _DATE_WEAK)), None)

    name_col = next((i for i, h in enumerate(headers)
                     if any(k in h for k in _NAME_KEYS)), None)

    data_start = 1
    if date_col is None:
        # Possibly a headerless table — only accept it if the FIRST cell of the
        # first row actually parses as a date. Otherwise it isn't a holiday table.
        year_match = re.search(r'20\d{2}', doc_id)
        default_year = int(year_match.group()) if year_match else 2026
        if len(rows[0]) >= 2 and normalize_date(clean_table_cell(rows[0][0]), default_year):
            date_col, name_col, data_start = 0, 1, 0
        else:
            return []

    # Skip "restricted / optional holiday" tables — office stays open on those.
    if name_col is not None and name_col < len(headers):
        if any(k in headers[name_col] for k in _NON_WORKING_HEADER):
            logger.info("Skipping non-working-day holiday table", extra={"doc_id": doc_id,
                                                                        "header": headers[name_col]})
            return []

    if name_col is None or name_col == date_col:
        # Fall back to the last column as the holiday name
        name_col = len(headers) - 1 if len(headers) > 1 else None
        if name_col == date_col:
            return []

    # Determine default year from doc_id or use 2026
    year_match = re.search(r'20\d{2}', doc_id)
    default_year = int(year_match.group()) if year_match else 2026

    holidays = []
    for row in rows[data_start:]:
        if not any(row):
            continue
        date_str = clean_table_cell(row[date_col]) if date_col < len(row) else ""
        name_str = clean_table_cell(row[name_col]) if name_col is not None and name_col < len(row) else ""

        if not date_str or not name_str:
            continue

        iso_date = normalize_date(date_str, default_year)
        if not iso_date:
            logger.debug("Skipping non-date row in holiday table",
                         extra={"raw": date_str, "doc_id": doc_id})
            continue

        holidays.append({
            "region": region,
            "date": iso_date,
            "name": name_str,
            "year": int(iso_date[:4]),
            "doc_id": doc_id,
        })

    return holidays


# ─── Inserter ─────────────────────────────────────────────────────────────────

def upsert_holidays(holidays: list[dict]) -> int:
    """Insert or replace holiday records. Returns count inserted."""
    if not holidays:
        return 0
    ensure_holidays_table()
    conn = get_db()
    count = 0
    for h in holidays:
        try:
            conn.execute(
                """
                INSERT OR REPLACE INTO holidays (region, date, name, year, doc_id)
                VALUES (?, ?, ?, ?, ?)
                """,
                (h["region"], h["date"], h["name"], h["year"], h["doc_id"]),
            )
            count += 1
        except Exception as exc:
            logger.warning("Holiday insert failed", extra={"date": h.get("date"), "error": str(exc)})
    conn.commit()
    logger.info("Holidays upserted", extra={"count": count})
    return count


# ─── Entry point for each document ───────────────────────────────────────────

# Type alias: a "table" is a list of rows; a document yields a list of tables.
Table = list[list[str]]


def extract_holidays_from_document(file_path: str, doc_meta: dict) -> int:
    """
    Main entry point. Extracts holidays from a calendar document.

    Every file yields a list of tables; each table is parsed independently so
    document-metadata tables are ignored and only genuine holiday tables are used.
    Returns the number of holidays inserted.
    """
    doc_id = doc_meta["doc_id"]
    regions = doc_meta.get("regions", [])
    region = regions[0] if regions else "GLOBAL"

    ext = file_path.lower().split(".")[-1]

    if ext == "pdf":
        tables = _extract_tables_from_pdf(file_path)
    elif ext in ("doc", "docx"):
        tables = _extract_tables_from_docx(file_path)
    elif ext in ("html", "htm"):
        tables = _extract_tables_from_html(file_path)
    elif ext == "md":
        tables = _extract_tables_from_markdown(file_path)
    else:
        tables = []

    if not tables:
        logger.warning("No holiday tables extracted", extra={"file": file_path})
        return 0

    holidays: list[dict] = []
    for rows in tables:
        holidays.extend(parse_holiday_table(rows, doc_id, region))

    # De-duplicate by (region, date) — later/duplicate rows overwrite earlier ones.
    deduped: dict[tuple[str, str], dict] = {(h["region"], h["date"]): h for h in holidays}

    logger.info("Holiday tables parsed",
                extra={"file": os.path.basename(file_path), "tables": len(tables),
                       "holidays": len(deduped)})
    return upsert_holidays(list(deduped.values()))


def _extract_tables_from_pdf(file_path: str) -> list[Table]:
    """Extract every table from a PDF using PyMuPDF's table finder."""
    try:
        import fitz
    except ImportError:
        logger.error("PyMuPDF not installed — cannot extract PDF tables")
        return []

    tables: list[Table] = []
    try:
        pdf = fitz.open(file_path)
        for page in pdf:
            try:
                found = page.find_tables()
            except Exception:
                found = None
            if not found:
                continue
            for table in found.tables:
                rows = [[clean_table_cell(c) for c in row] for row in table.extract()]
                rows = [r for r in rows if any(r)]
                if rows:
                    tables.append(rows)
        pdf.close()
    except Exception as exc:
        logger.warning("PDF holiday extract error", extra={"file": file_path, "error": str(exc)})
    return tables


def _extract_tables_from_docx(file_path: str) -> list[Table]:
    """Return each DOCX table as a list of rows."""
    try:
        from docx import Document
        doc = Document(file_path)
        tables: list[Table] = []
        for table in doc.tables:
            rows = [[clean_table_cell(c.text) for c in row.cells] for row in table.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                tables.append(rows)
        return tables
    except Exception as exc:
        logger.warning("DOCX holiday extract error", extra={"file": file_path, "error": str(exc)})
        return []


def _extract_tables_from_html(file_path: str) -> list[Table]:
    """Return each HTML <table> as a list of rows."""
    try:
        from bs4 import BeautifulSoup
        with open(file_path, encoding="utf-8", errors="replace") as f:
            soup = BeautifulSoup(f.read(), "lxml")
        tables: list[Table] = []
        for table_tag in soup.find_all("table"):
            rows = []
            for tr in table_tag.find_all("tr"):
                cells = [clean_table_cell(td.get_text()) for td in tr.find_all(["td", "th"])]
                if any(cells):
                    rows.append(cells)
            if rows:
                tables.append(rows)
        return tables
    except Exception as exc:
        logger.warning("HTML holiday extract error", extra={"file": file_path, "error": str(exc)})
        return []


def _extract_tables_from_markdown(file_path: str) -> list[Table]:
    """Parse markdown pipe-tables into a list of tables."""
    try:
        with open(file_path, encoding="utf-8", errors="replace") as f:
            content = f.read()

        def _is_separator(cells: list[str]) -> bool:
            return all(c and set(c) <= set("-: ") for c in cells)

        tables: list[Table] = []
        current: Table = []
        for line in content.split("\n"):
            if "|" in line:
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if _is_separator(cells):
                    continue
                if any(cells):
                    current.append(cells)
            else:
                if current:
                    tables.append(current)
                    current = []
        if current:
            tables.append(current)
        return tables
    except Exception as exc:
        logger.warning("Markdown holiday extract error", extra={"file": file_path, "error": str(exc)})
        return []
